#!/usr/bin/env python3
"""
svg_to_pptx.py - rebuild a matplotlib SVG as native PowerPoint shapes

Purpose
    Converts a figure saved by matplotlib as SVG into a one-slide .pptx in
    which every curve, filled area, marker, tick and label is a separate,
    editable PowerPoint shape. Rasterised artists (for instance a contact map
    drawn with `rasterized=True`) are placed as pictures at their position.

Inputs
    svg     an SVG written by matplotlib. Text is rebuilt as text boxes only
            when the figure was saved with rcParams["svg.fonttype"] = "none";
            text drawn as glyph outlines is skipped with a message.

Outputs
    pptx    one slide of the size of the SVG, in points

Method
    The slide keeps the point coordinates of the SVG, so each element lands
    where matplotlib put it. Paths made of M/L/z commands become freeform
    shapes with their fill, stroke, width and opacity; Bezier segments (Q, C)
    are flattened into --curve-steps chords. A path clipped to an
    axes box is clipped to its own clip rectangle (Sutherland-Hodgman for a
    closed polygon, clamping for an open line). Markers, ticks and the
    polygons of fill_between are written by matplotlib as <use> of a path
    defined once; each <use> becomes a shape offset to its x, y, styled by
    the <use> itself when it carries a style. Embedded PNG images are
    decoded, flipped as their transform requires, and inserted as pictures.
    White fills without stroke (figure and axes backgrounds, label boxes) are
    left out unless --keep-white is given. Shapes are added in document
    order, so the stacking order of the SVG is preserved.

Usage
    python3 svg_to_pptx.py figures/locus_GENE.svg figures/locus_GENE.pptx
    python3 svg_to_pptx.py figures/folding_GENE.svg figures/folding_GENE.pptx \\
        --font "Helvetica"

References
    Hunter JD. 2007. Matplotlib: a 2D graphics environment. Computing in
        Science and Engineering.
    python-pptx (software). https://github.com/scanny/python-pptx
"""

from __future__ import annotations

import argparse
import base64
import io
import re
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

SVG_NS = "http://www.w3.org/2000/svg"
XLINK_HREF = "{http://www.w3.org/1999/xlink}href"
NS = {"svg": SVG_NS}
NUM = re.compile(r"-?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?")
# groups whose paths, markers and texts are converted
DRAWN_GROUPS = ("FillBetweenPolyCollection", "PolyCollection", "patch",
                "LineCollection", "line2d", "text", "PathCollection")


def style(el):
    d = {}
    for part in (el.get("style") or "").split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            d[k.strip()] = v.strip()
    return d


def path_points(d, steps=16):
    """Vertices of a path with absolute M, L, Q, C and z commands, in SVG
    units. Quadratic and cubic Bezier segments are flattened into `steps`
    chords each, so curves keep their shape as freeform outlines."""
    pts, cmd, cur = [], "M", []
    for tok in re.findall(r"[MLQCzZ]|-?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?", d):
        if tok in "MLQCzZ":
            cmd, cur = tok, []
            continue
        cur.append(float(tok))
        need = {"Q": 4, "C": 6}.get(cmd, 2)
        if len(cur) < need:
            continue
        if cmd in "QC" and pts:
            ctrl = [pts[-1]] + [(cur[k], cur[k + 1])
                                for k in range(0, need, 2)]
            for s in range(1, steps + 1):
                t = s / steps
                if cmd == "Q":
                    w = ((1 - t) ** 2, 2 * (1 - t) * t, t ** 2)
                else:
                    w = ((1 - t) ** 3, 3 * (1 - t) ** 2 * t,
                         3 * (1 - t) * t ** 2, t ** 3)
                pts.append((sum(wi * p[0] for wi, p in zip(w, ctrl)),
                            sum(wi * p[1] for wi, p in zip(w, ctrl))))
        else:
            pts.append((cur[-2], cur[-1]))
        cur = []
    return pts


def parse_transform(tr):
    """An SVG transform list as an affine (a, b, c, d, e, f)."""
    def mul(m, n):
        a, b, c, d, e, f = m
        a2, b2, c2, d2, e2, f2 = n
        return (a * a2 + c * b2, b * a2 + d * b2,
                a * c2 + c * d2, b * c2 + d * d2,
                a * e2 + c * f2 + e, b * e2 + d * f2 + f)
    m = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)
    for name, args in re.findall(r"(\w+)\s*\(([^)]*)\)", tr or ""):
        v = [float(x) for x in NUM.findall(args)]
        if name == "matrix" and len(v) == 6:
            n = tuple(v)
        elif name == "translate":
            n = (1.0, 0.0, 0.0, 1.0, v[0], v[1] if len(v) > 1 else 0.0)
        elif name == "scale":
            n = (v[0], 0.0, 0.0, v[1] if len(v) > 1 else v[0], 0.0, 0.0)
        else:
            continue
        m = mul(m, n)
    return m


def clip_box(pts, x0, y0, x1, y1, closed):
    """Sutherland-Hodgman for a closed polygon, plain clamping for a line."""
    if not closed:
        return [(min(max(x, x0), x1), min(max(y, y0), y1)) for x, y in pts]

    def inside(p, edge):
        x, y = p
        return (x >= x0, x <= x1, y >= y0, y <= y1)[edge]

    def cut(a, b, edge):
        (ax, ay), (bx, by) = a, b
        if edge < 2:
            xe = x0 if edge == 0 else x1
            t = (xe - ax) / (bx - ax) if bx != ax else 0.0
            return (xe, ay + t * (by - ay))
        ye = y0 if edge == 2 else y1
        t = (ye - ay) / (by - ay) if by != ay else 0.0
        return (ax + t * (bx - ax), ye)

    out = list(pts)
    for edge in range(4):
        if not out:
            break
        src, out = out, []
        for i, b in enumerate(src):
            a = src[i - 1]
            ia, ib = inside(a, edge), inside(b, edge)
            if ib:
                if not ia:
                    out.append(cut(a, b, edge))
                out.append(b)
            elif ia:
                out.append(cut(a, b, edge))
    return out


def dedup(pts, eps=0.02):
    keep = []
    for p in pts:
        if not keep or abs(p[0] - keep[-1][0]) > eps or \
                abs(p[1] - keep[-1][1]) > eps:
            keep.append(p)
    return keep


def colour(value):
    """An SVG colour as RRGGBB, or None for none / unsupported values."""
    if not value or value == "none":
        return None
    value = value.strip()
    if re.fullmatch(r"#[0-9a-fA-F]{6}", value):
        return value[1:]
    if re.fullmatch(r"#[0-9a-fA-F]{3}", value):
        return "".join(ch * 2 for ch in value[1:])
    return {"black": "000000", "white": "FFFFFF"}.get(value.lower())


def drop_theme_style(sh):
    """The default theme style reference carries a drop shadow."""
    sh.shadow.inherit = False
    st_el = sh._element.find(qn("p:style"))
    if st_el is not None:
        sh._element.remove(st_el)


def add_poly(slide, pts, fill=None, line=None, width=0.8, alpha=None,
             closed=True):
    pts = dedup(pts)
    if len(pts) < 2:
        return None
    b = slide.shapes.build_freeform(Pt(pts[0][0]), Pt(pts[0][1]))
    b.add_line_segments([(Pt(x), Pt(y)) for x, y in pts[1:]], close=closed)
    sh = b.convert_to_shape()
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        if alpha is not None:
            el = sh.fill.fore_color._xFill.find(qn("a:srgbClr"))
            a = el.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
            el.append(a)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(width)
    else:
        sh.line.fill.background()
    drop_theme_style(sh)
    return sh


def add_text(slide, txt, x, y, size=9, anchor="start", rotate=0, font="Arial",
             rgb="000000", italic=False, bold=False):
    """x, y are the SVG anchor point; y is the text baseline."""
    w = max(6.0, 0.62 * size * len(txt) + 4)
    h = 1.7 * size
    if rotate:
        left, top = x - w / 2, y - h / 2
    else:
        left = {"start": x, "middle": x - w / 2, "end": x - w}[anchor]
        top = y - 1.12 * size
    box = slide.shapes.add_textbox(Pt(left), Pt(top), Pt(w), Pt(h))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER,
                   "end": PP_ALIGN.RIGHT}[anchor]
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(size)
    r.font.name = font
    r.font.color.rgb = RGBColor.from_string(rgb)
    r.font.italic = italic or None
    r.font.bold = bold or None
    if rotate:
        box.rotation = (360 + rotate) % 360
    box.shadow.inherit = False
    return box


def add_image(slide, img):
    """An embedded PNG, placed through its transform."""
    href = img.get(XLINK_HREF) or img.get("href") or ""
    if not href.startswith("data:image/png;base64,"):
        return False
    from PIL import Image
    data = base64.b64decode(href.split(",", 1)[1])
    x, y = float(img.get("x", 0)), float(img.get("y", 0))
    w, h = float(img.get("width")), float(img.get("height"))
    a, b, c, d, e, f = parse_transform(img.get("transform"))
    corners = [(a * px + c * py + e, b * px + d * py + f)
               for px, py in ((x, y), (x + w, y), (x, y + h), (x + w, y + h))]
    xs, ys = [p[0] for p in corners], [p[1] for p in corners]
    pic = Image.open(io.BytesIO(data))
    # a negative scale mirrors the stored bitmap
    if a < 0:
        pic = pic.transpose(Image.Transpose.FLIP_LEFT_RIGHT)
    if d < 0:
        pic = pic.transpose(Image.Transpose.FLIP_TOP_BOTTOM)
    buf = io.BytesIO()
    pic.save(buf, format="PNG")
    buf.seek(0)
    slide.shapes.add_picture(buf, Pt(min(xs)), Pt(min(ys)),
                             Pt(max(xs) - min(xs)), Pt(max(ys) - min(ys)))
    return True


def main():
    p = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("svg", help="SVG written by matplotlib")
    p.add_argument("pptx", help="output .pptx")
    p.add_argument("--font", default="Arial",
                   help="font given to the rebuilt text boxes")
    p.add_argument("--keep-white", action="store_true",
                   help="keep white filled shapes without stroke (figure and "
                        "axes backgrounds, label boxes)")
    p.add_argument("--curve-steps", type=int, default=16,
                   help="chords per Bezier segment when curves are flattened")
    p.add_argument("--no-images", action="store_true",
                   help="leave out the embedded raster images")
    a = p.parse_args()

    root = ET.parse(a.svg).getroot()
    W = float(NUM.search(root.get("width")).group())
    H = float(NUM.search(root.get("height")).group())

    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W), Pt(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # each axes has its own clip rectangle, referenced as url(#id)
    clips = {}
    for cp in root.iter(f"{{{SVG_NS}}}clipPath"):
        r = cp.find("svg:rect", NS)
        if r is not None:
            x0, y0 = float(r.get("x")), float(r.get("y"))
            clips[cp.get("id")] = (x0, y0, x0 + float(r.get("width")),
                                   y0 + float(r.get("height")))

    def clip_of(el):
        m = re.search(r"url\(#([^)]+)\)", el.get("clip-path") or "")
        return clips.get(m.group(1)) if m else None

    n_fill = n_line = n_txt = n_img = n_glyph = 0
    for g in root.iter(f"{{{SVG_NS}}}g"):
        if not a.no_images:
            for img in g.findall("svg:image", NS):
                n_img += add_image(slide, img)

        gid = g.get("id") or ""
        if not gid.startswith(DRAWN_GROUPS):
            continue

        for pth in g.findall("svg:path", NS):
            st = style(pth)
            pts = path_points(pth.get("d"), a.curve_steps)
            if len(pts) < 2:
                continue
            fill = colour(st.get("fill", "none"))
            stroke = colour(st.get("stroke", "none"))
            if fill is None and stroke is None:
                continue
            if fill and fill.upper() == "FFFFFF" and stroke is None \
                    and not a.keep_white:
                continue
            closed = pth.get("d").rstrip().endswith(("z", "Z"))
            box = clip_of(pth)
            if box:
                pts = clip_box(pts, *box, closed)
            if len(pts) < 2:
                continue
            op = st.get("fill-opacity") or st.get("stroke-opacity")
            add_poly(slide, pts, fill=fill, line=stroke,
                     width=float(st.get("stroke-width", 0.8)),
                     alpha=float(op) if op else None, closed=closed)
            if fill:
                n_fill += 1
            else:
                n_line += 1

        # markers, ticks and fill_between polygons: <use> of a defined path
        for use in g.findall(".//svg:use", NS):
            if use.get("x") is None or use.get("y") is None:
                n_glyph += 1          # a glyph of text drawn as outlines
                continue
            ref = (use.get(XLINK_HREF) or use.get("href") or "").lstrip("#")
            shape = root.find(f".//svg:path[@id='{ref}']", NS)
            if shape is None:
                continue
            x, y = float(use.get("x")), float(use.get("y"))
            pts = [(x + dx, y + dy)
                   for dx, dy in path_points(shape.get("d"), a.curve_steps)]
            st = style(use) or style(shape)
            fill = colour(st.get("fill", "none"))
            stroke = colour(st.get("stroke", "none"))
            if fill is None and stroke is None:
                stroke = "000000"
            closed = shape.get("d").rstrip().endswith(("z", "Z"))
            op = st.get("fill-opacity")
            add_poly(slide, pts, fill=fill, line=stroke,
                     width=float(st.get("stroke-width",
                                        style(shape).get("stroke-width", 0.8))),
                     alpha=float(op) if op else None, closed=closed)
            if fill:
                n_fill += 1
            else:
                n_line += 1

        for t in g.findall("svg:text", NS):
            st = style(t)
            size = float(NUM.search(st.get("font-size", "9px")).group())
            anchor = {"end": "end", "middle": "middle"}.get(
                st.get("text-anchor", "start"), "start")
            tr = t.get("transform") or ""
            rot = 0
            if "rotate" in tr:
                rot = float(NUM.search(tr[tr.index("rotate"):]).group())
            if "translate" in tr:
                v = NUM.findall(tr[tr.index("translate"):])
                x, y = float(v[0]), float(v[1])
            else:
                x, y = float(t.get("x", 0)), float(t.get("y", 0))
            add_text(slide, "".join(t.itertext()), x, y, size, anchor, rot,
                     font=a.font, rgb=colour(st.get("fill")) or "000000",
                     italic=st.get("font-style") == "italic",
                     bold=st.get("font-weight") in ("bold", "700", "800",
                                                    "900"))
            n_txt += 1

    prs.save(a.pptx)
    print(f"{a.pptx}: {n_fill} filled shapes, {n_line} lines, {n_txt} text "
          f"boxes, {n_img} images, slide {W/72:.2f} x {H/72:.2f} in")
    if n_glyph:
        print(f"  {n_glyph} glyphs drawn as outlines were skipped; save the "
              f"figure with rcParams['svg.fonttype'] = 'none' to keep its "
              f"text editable")


if __name__ == "__main__":
    main()
